"""Read-only PPTX content audit using ZIP/XML and embedded Excel data.

Example (repository root):
    python tools/verify_defense_content.py --pptx tmp/presentations/defense/20261002_214023/candidate.pptx

The deck is never opened in PowerPoint, rebuilt, extracted, or written. Only the
requested JSON audit is written. This verifies content; rendering is separate.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import posixpath
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parents[1]
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
PENDING = re.compile(
    r"\b(?:TBD|TODO)\b|\b(?:results?|outcomes?)\s+(?:are\s+)?(?:pending|in\s+progress)\b"
    r"|\bpending\s+(?:results?|outcomes?)\b|\bexperiment\s+is\s+still\s+in\s+progress\b",
    re.IGNORECASE,
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(value: Any) -> str:
    return " ".join(str(value).split())


def text(node: ET.Element | None) -> str:
    if node is None:
        return ""
    paragraphs = node.findall(".//a:p", NS)
    return "\n".join("".join(p.itertext()) if not p.findall(".//a:t", NS)
                     else "".join(t.text or "" for t in p.findall(".//a:t", NS))
                     for p in paragraphs)


def xml(archive: ZipFile, name: str) -> ET.Element:
    return ET.fromstring(archive.read(name))


def relationships(archive: ZipFile, owner: str) -> dict[str, dict[str, str]]:
    directory, filename = posixpath.split(owner)
    part = posixpath.join(directory, "_rels", filename + ".rels")
    if part not in archive.namelist():
        return {}
    result = {}
    for relationship in xml(archive, part).findall("rel:Relationship", NS):
        target = relationship.attrib["Target"]
        result[relationship.attrib["Id"]] = dict(relationship.attrib) | {
            "resolved": posixpath.normpath(posixpath.join(directory, target))
        }
    return result


def table_cells(node: ET.Element) -> list[list[str]]:
    return [[normalize(text(cell)) for cell in row.findall("a:tc", NS)]
            for row in node.findall("a:tr", NS)]


def cache_points(parent: ET.Element | None, numeric: bool) -> tuple[list[Any], str]:
    if parent is None:
        raise ValueError("Missing chart data reference")
    reference = parent.find("c:numRef" if numeric else "c:strRef", NS)
    if reference is None:
        raise ValueError("Chart must reference an editable embedded worksheet")
    cache = reference.find("c:numCache" if numeric else "c:strCache", NS)
    if cache is None:
        raise ValueError("Missing chart cache")
    points = sorted(cache.findall("c:pt", NS), key=lambda p: int(p.attrib["idx"]))
    count = int(cache.find("c:ptCount", NS).attrib["val"])
    if count != len(points) or [int(p.attrib["idx"]) for p in points] != list(range(count)):
        raise ValueError("Chart point count or contiguous indices disagree")
    values = [p.findtext("c:v", default="", namespaces=NS) for p in points]
    if numeric:
        values = [float(v) for v in values]
        if not all(math.isfinite(v) for v in values):
            raise ValueError("Nonfinite chart cache value")
    return values, reference.findtext("c:f", default="", namespaces=NS)


def worksheet_reference(workbook: Any, formula: str) -> list[Any]:
    if "!" not in formula:
        raise ValueError(f"Unrecognized chart cell reference: {formula}")
    sheet, coordinates = formula.rsplit("!", 1)
    sheet = sheet.strip("'").replace("''", "'")
    if sheet not in workbook.sheetnames or "[" in sheet:
        raise ValueError(f"External or absent worksheet: {formula}")
    left, top, right, bottom = range_boundaries(coordinates)
    return [workbook[sheet].cell(row, column).value for row in range(top, bottom + 1)
            for column in range(left, right + 1)]


def close_values(actual: list[Any], expected: list[Any], tolerance: float) -> bool:
    return len(actual) == len(expected) and all(
        isinstance(a, (float, int)) and not isinstance(a, bool) and math.isfinite(a)
        and abs(float(a) - float(e)) <= tolerance for a, e in zip(actual, expected)
    )


def audit_chart(archive: ZipFile, part: str, specification: dict[str, Any],
                number: int, tolerance: float, require: Any) -> dict[str, Any]:
    chart = xml(archive, part)
    require(chart.find(".//c:barChart/c:barDir", NS) is not None and
            chart.find(".//c:barChart/c:barDir", NS).attrib.get("val") == "col",
            f"slide {number}: native column-chart kind")
    external = chart.find("c:externalData", NS)
    if external is None:
        raise ValueError("Chart has no embedded workbook relationship")
    rel = relationships(archive, part)[external.attrib[f"{{{NS['r']}}}id"]]
    if rel.get("TargetMode") == "External" or not rel["resolved"].startswith("ppt/embeddings/"):
        raise ValueError("Chart workbook is external")
    embedded = rel["resolved"]
    workbook = load_workbook(io.BytesIO(archive.read(embedded)), data_only=False, read_only=True)
    try:
        series_nodes = chart.findall(".//c:barChart/c:ser", NS)
        require(len(series_nodes) == len(specification["series"]), f"slide {number}: series count")
        series_results = []
        for index, expected in enumerate(specification["series"]):
            if index >= len(series_nodes):
                break
            series = series_nodes[index]
            categories, category_formula = cache_points(series.find("c:cat", NS), False)
            values, value_formula = cache_points(series.find("c:val", NS), True)
            names, name_formula = cache_points(series.find("c:tx", NS), False)
            worksheet_categories = worksheet_reference(workbook, category_formula)
            worksheet_values = worksheet_reference(workbook, value_formula)
            worksheet_names = worksheet_reference(workbook, name_formula)
            require(categories == specification["categories"], f"slide {number}, series {index}: cached categories")
            require(names == [expected["name"]], f"slide {number}, series {index}: cached series name")
            require(worksheet_categories == specification["categories"], f"slide {number}, series {index}: worksheet categories")
            require(worksheet_names == [expected["name"]], f"slide {number}, series {index}: worksheet series name")
            require(close_values(values, expected["values"], tolerance), f"slide {number}, series {index}: cached numerical values")
            require(close_values(worksheet_values, expected["values"], tolerance), f"slide {number}, series {index}: worksheet numerical values")
            require(close_values(values, worksheet_values, tolerance), f"slide {number}, series {index}: cache/worksheet identity")
            series_results.append({"name": expected["name"], "expected_values": expected["values"],
                "cached_values": values, "embedded_worksheet_values": worksheet_values,
                "maximum_absolute_cache_error": max((abs(a - b) for a, b in zip(values, expected["values"])), default=0),
                "formulas": {"name": name_formula, "categories": category_formula, "values": value_formula}})
        expected_matrix = [[specification["x_title"], *(s["name"] for s in specification["series"])]]
        expected_matrix.extend([[category, *(s["values"][i] for s in specification["series"])]
                                for i, category in enumerate(specification["categories"])])
        require(workbook.sheetnames == ["Sheet1"], f"slide {number}: one embedded data sheet")
        actual_matrix = [list(row) for row in workbook["Sheet1"].values]
        # Excel can retain formatted but empty slots from its default chart
        # sheet. Ignore only trailing empty rows/columns, never nonempty data.
        while actual_matrix and all(value is None for value in actual_matrix[-1]):
            actual_matrix.pop()
        used_width = max((column + 1 for row in actual_matrix for column, value in enumerate(row)
                          if value is not None), default=0)
        actual_matrix = [row[:used_width] for row in actual_matrix]
        matrix_matches = len(actual_matrix) == len(expected_matrix)
        for actual_row, expected_row in zip(actual_matrix, expected_matrix):
            matrix_matches &= len(actual_row) == len(expected_row)
            for actual_value, expected_value in zip(actual_row, expected_row):
                if isinstance(expected_value, (float, int)):
                    matrix_matches &= close_values([actual_value], [expected_value], tolerance)
                else:
                    matrix_matches &= actual_value == expected_value
        require(matrix_matches, f"slide {number}: full embedded worksheet (no stale sample data)")
    finally:
        workbook.close()
    category_axis = chart.find(".//c:catAx", NS)
    value_axis = chart.find(".//c:valAx", NS)
    require(normalize(text(category_axis.find("c:title", NS))) == normalize(specification["x_title"]),
            f"slide {number}: category axis title")
    require(normalize(text(value_axis.find("c:title", NS))) == normalize(specification["y_axis_title"]),
            f"slide {number}: value axis title")
    for bound in ("min", "max"):
        expected = specification.get(f"y_axis_{bound}")
        if expected is not None:
            actual = value_axis.find(f"c:scaling/c:{bound}", NS)
            require(actual is not None and abs(float(actual.attrib["val"]) - expected) <= tolerance,
                    f"slide {number}: value-axis {bound}")
    return {"slide_number": number, "chart_part": part, "embedded_workbook": embedded,
            "categories": specification["categories"], "series": series_results,
            "embedded_matrix": actual_matrix}


def appendix_d_review(require: Any) -> dict[str, Any]:
    """Independently invert binomial tails rather than reuse the tool's beta PPF."""
    from scipy.optimize import brentq
    from scipy.stats import binom

    appendix = ROOT / "ThesisDocs/appendices.md"
    tool = ROOT / "tools/analyze_live_stopping_uncertainty.py"
    paragraph = appendix.read_text(encoding="utf-8").split("# Appendix D Paired accuracy uncertainty", 1)[1]
    require("97.5%" in paragraph and "0.025" in paragraph and "0.0125" in paragraph,
            "Appendix D: declared marginal/tail errors match 97.5% two-sided CP")
    require("independent and identically distributed" in paragraph and "does not assume their independence" in paragraph,
            "Appendix D: iid across task pairs, no within-pair independence assumption")
    require("hand-selected" in paragraph and "does not provide randomized coverage" in paragraph,
            "Appendix D: fixed challenge bank reference-model qualification")

    def cp(k: int, n: int) -> tuple[float, float]:
        tail = .0125
        lower = 0.0 if k == 0 else brentq(lambda p: binom.sf(k - 1, n, p) - tail, 0, 1, xtol=1e-14)
        upper = 1.0 if k == n else brentq(lambda p: binom.cdf(k, n, p) - tail, 0, 1, xtol=1e-14)
        return lower, upper

    base = ROOT / "research/outputs/semester2/online_stopping_20261002"
    intervals = []
    for name in ("", "adversarial_live", "learned_main", "learned_adversarial"):
        folder = base / name
        recorded = json.loads((folder / "live_uncertainty.json").read_text(encoding="utf-8"))
        counts = recorded["counts"]
        n = counts["problems_or_trajectories"]
        plus, minus = cp(counts["paired_improved"], n), cp(counts["paired_worsened"], n)
        interval = [plus[0] - minus[1], plus[1] - minus[0]]
        require(close_values(interval, recorded["accuracy_delta_conservative_exact_95ci"], 1e-11),
                f"Appendix D: independent binomial inversion matches {name or 'heuristic_main'}")
        if counts["paired_improved"] == counts["paired_worsened"] == 0:
            closed = 1 - .0125 ** (1 / n)
            require(close_values(interval, [-closed, closed], 1e-11), f"Appendix D: zero-discordance formula n={n}")
        intervals.append({"folder": str(folder.relative_to(ROOT)), "n": n,
            "improved": counts["paired_improved"], "worsened": counts["paired_worsened"],
            "independent_accuracy_delta_interval": interval,
            "recorded_accuracy_delta_interval": recorded["accuracy_delta_conservative_exact_95ci"]})
    return {"appendix_sha256": sha256(appendix), "uncertainty_tool_sha256": sha256(tool),
        "proof_assessment": "Correct conditional on the stated iid task-pair law: each marginal is Binomial(n,pi); two CP intervals have marginal error<=.025; a union bound gives joint error<=.05 without independence between discordance categories; interval subtraction bounds delta.",
        "assumption_limits": [
            "Unique task IDs do not establish independent, identically distributed sampling. The tool checks uniqueness/binary outcomes, not the scientific sampling law.",
            "The shuffled cached GSM8K panel is selected without replacement; the iid CP interval is a declared reference-model calculation, not exact finite-population design coverage.",
            "The handpicked trap bank has no randomized adversarial-population coverage; no adaptive policy selection, cross-model generalization, or noninferiority guarantee is covered.",
            "Paired token bootstrap resamples whole problem pairs and computes 1-sum(active)/sum(baseline); its percentile interval is descriptive, not an exact finite-sample certificate.",
        ], "intervals": intervals}


def scientific_chart_sources(source_slides: list[dict[str, Any]], tolerance: float, require: Any) -> list[dict[str, Any]]:
    """Check the JSON chart inputs against their declared scientific CSVs."""
    results = []
    for number in (15, 17, 18):
        slide = source_slides[number - 1]
        path = ROOT / next(name for name in slide["sources"] if name.endswith(".csv"))
        with path.open(encoding="utf-8", newline="") as handle:
            rows = list(csv.DictReader(handle))
        if number == 15:
            keys = [("qwen2p5_7b__gsm8k", "4"), ("qwen2p5_7b__gsm8k", "5"),
                    ("qwen2p5_32b__math", "5"), ("qwen2p5_32b__math", "6")]
            values = []
            for cell, step in keys:
                matches = [row for row in rows if row["cell"] == cell and row["step"] == step]
                if len(matches) != 1:
                    raise ValueError(f"Scientific chart source has nonunique cell/step: {cell}/{step}")
                values.append(float(matches[0]["net_drift"]))
            selection = {"cell_step_keys": keys, "column": "net_drift"}
        elif number == 17:
            precision = [row for row in rows if row["experiment"] == "N6"]
            if len(precision) != 1:
                raise ValueError("Scientific chart source has no unique N6 precision row")
            counts = re.search(r"(\d+)/(\d+) BF16 correct versus (\d+)/(\d+) 4-bit correct",
                               precision[0]["raw_aggregate"])
            if not counts:
                raise ValueError("Precision source omitted arm correctness counts")
            bf_correct, bf_n, quant_correct, quant_n = map(int, counts.groups())
            values = [bf_correct / bf_n, quant_correct / quant_n]
            selection = {"experiment": "N6", "count_numerators": [bf_correct, quant_correct],
                         "count_denominators": [bf_n, quant_n]}
        else:
            gru = [row for row in rows if row["configuration"] == "Causal GRU"]
            if len(gru) != 1:
                raise ValueError("Scientific chart source has no unique causal GRU row")
            columns = ["micro_oof_auc", "task_macro_auc", "domain_macro_auc", "worst_domain_auc"]
            values = [float(gru[0][column]) for column in columns]
            selection = {"configuration": "Causal GRU", "columns": columns}
            for row in slide.get("table_secondary", {}).get("rows", []):
                scientific = [r for r in rows if r["configuration"] == row[0]]
                require(len(scientific) == 1, f"slide {number}: secondary table configuration in scientific source")
                if scientific:
                    require(row[1:] == [f"{float(scientific[0]['micro_oof_auc']):.4f}",
                                        f"{float(scientific[0]['micro_step_utility']):.4f}"],
                            f"slide {number}: secondary table rounded AUC/utility values match scientific source")
        require(close_values(slide["chart"]["series"][0]["values"], values, tolerance),
                f"slide {number}: numerical JSON inputs match scientific CSV")
        results.append({"slide_number": number, "scientific_source": str(path.relative_to(ROOT)),
                        "source_sha256": sha256(path), "selection": selection, "derived_values": values})
    return results


def audit(pptx: Path, source_path: Path, tolerance: float) -> dict[str, Any]:
    source = json.loads(source_path.read_text(encoding="utf-8"))
    source_slides = source["slides"]
    checks, defects = {}, []

    def require(condition: bool, name: str) -> None:
        checks[name] = bool(condition)
        if not condition:
            defects.append(name)

    before = sha256(pptx)
    require(source["slide_count"] == len(source_slides) == 25, "source declares 25 slides")
    require([s["number"] for s in source_slides] == list(range(1, 26)), "source slide numbers consecutive")
    require(sum(s["seconds"] for s in source_slides) == source["total_seconds"] == 1800, "source durations sum to 1800 seconds")
    expected_tables = sum(("table" in s) + ("table_secondary" in s) for s in source_slides)
    require(sum("table" in s for s in source_slides) == 12, "source declares 12 primary native tables")
    require(sum("chart" in s for s in source_slides) == 3, "source declares 3 native charts")
    require(not any(s.get("pending_outcomes") for s in source_slides), "source has no pending outcomes")
    slides, charts, tables, pending = [], [], [], []
    notes_total, actual_table_count = 0, 0
    chart_parts_seen, notes_parts_seen = set(), set()
    with ZipFile(pptx) as archive:
        require(archive.testzip() is None, "PPTX ZIP CRC integrity")
        presentation = xml(archive, "ppt/presentation.xml")
        relations = relationships(archive, "ppt/presentation.xml")
        slide_parts = [relations[item.attrib[f"{{{NS['r']}}}id"]]["resolved"]
                       for item in presentation.findall("p:sldIdLst/p:sldId", NS)]
        require(len(slide_parts) == 25, "PPTX has 25 ordered slides")
        for specification, part in zip(source_slides, slide_parts):
            number = specification["number"]
            slide = xml(archive, part)
            all_text = text(slide)
            require(normalize(specification["title"]) in normalize(all_text), f"slide {number}: title matches source")
            native_tables = slide.findall(".//a:tbl", NS)
            wanted_tables = [(key, specification[key]) for key in ("table", "table_secondary") if key in specification]
            require(len(native_tables) == len(wanted_tables), f"slide {number}: native table count")
            actual_table_count += len(native_tables)
            for table_node, (key, expected) in zip(native_tables, wanted_tables):
                actual = table_cells(table_node)
                wanted = [[normalize(cell) for cell in row] for row in [expected["headers"], *expected["rows"]]]
                require(actual == wanted, f"slide {number}: {key} all cells match")
                tables.append({"slide_number": number, "source_key": key, "rows": actual,
                               "source_matches": actual == wanted})
            rel = relationships(archive, part)
            native_charts = slide.findall(".//c:chart", NS)
            require(len(native_charts) == (1 if "chart" in specification else 0), f"slide {number}: native chart count")
            for chart_ref in native_charts:
                chart_part = rel[chart_ref.attrib[f"{{{NS['r']}}}id"]]["resolved"]
                chart_parts_seen.add(chart_part)
                try:
                    charts.append(audit_chart(archive, chart_part, specification["chart"], number, tolerance, require))
                except Exception as error:
                    require(False, f"slide {number}: chart inspection failed: {type(error).__name__}: {error}")
            note_parts = [r["resolved"] for r in rel.values() if r["Type"].endswith("/notesSlide")]
            require(len(note_parts) == 1, f"slide {number}: one linked notes part")
            if not note_parts:
                continue
            note_part = note_parts[0]
            notes_parts_seen.add(note_part)
            note_root = xml(archive, note_part)
            body = next((sp for sp in note_root.findall(".//p:sp", NS)
                         if sp.find("p:nvSpPr/p:nvPr/p:ph", NS) is not None
                         and sp.find("p:nvSpPr/p:nvPr/p:ph", NS).attrib.get("type") == "body"), None)
            note_text = text(body)
            durations = re.findall(r"Target duration:\s*(\d+)\s*seconds\.", note_text)
            require(durations == [str(specification["seconds"])], f"slide {number}: notes duration matches source")
            duration = int(durations[0]) if len(durations) == 1 else 0
            notes_total += duration
            require(normalize(specification["speaker_notes"]) in normalize(note_text), f"slide {number}: full speaker notes present")
            references = [*specification.get("sources", []),
                          *(f"{reference['title']}: {reference['url']}" if isinstance(reference, dict) else reference
                            for reference in specification.get("references", []))]
            require(all(normalize(reference) in normalize(note_text) for reference in references), f"slide {number}: source references in notes")
            for scope, contents in (("slide", all_text), ("notes", note_text), ("source", json.dumps(specification))):
                for match in PENDING.finditer(contents):
                    pending.append({"slide_number": number, "scope": scope, "text": match.group()})
            slides.append({"number": number, "slide_part": part, "notes_part": note_part,
                           "title": specification["title"], "notes_duration_seconds": duration})
        package_charts = {n for n in archive.namelist() if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)}
        package_notes = {n for n in archive.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)}
        require(package_charts == chart_parts_seen and len(package_charts) == 3, "all 3 chart parts are linked to expected slides")
        require(package_notes == notes_parts_seen and len(package_notes) == 25, "all 25 notes parts are linked once")
        require(actual_table_count == expected_tables, f"PPTX has all {expected_tables} native editable source tables")
        require(notes_total == 1800, "25 notes durations sum to 1800 seconds")
        require(not pending, "no pending result text in slide/source/notes")
    scientific_sources = scientific_chart_sources(source_slides, tolerance, require)
    proof = appendix_d_review(require)
    require(sha256(pptx) == before, "PPTX bytes unchanged by read-only audit")
    return {"schema": "independent-native-defense-content-audit-v1", "made_at_utc": datetime.now(timezone.utc).isoformat(),
            "pptx": str(pptx), "pptx_sha256": before, "source": str(source_path), "source_sha256": sha256(source_path),
            "checker_sha256": sha256(Path(__file__)), "method": "ZIP/XML + openpyxl read-only; no COM, PowerPoint opening, rendering or rebuild",
            "numeric_absolute_tolerance": tolerance, "slide_count": len(slides), "native_table_count": actual_table_count,
            "expected_native_table_count": expected_tables,
            "native_chart_count": len(charts), "notes_duration_seconds": notes_total, "pending_result_text": pending,
            "charts": charts, "tables": tables, "slides": slides, "scientific_chart_sources": scientific_sources,
            "appendix_d_review": proof,
            "checks": checks, "defects": defects, "all_passed": all(checks.values()),
            "scope": "Native content/editability and numerical identity only; PowerPoint reopen and visual fit remain separate checks."}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--source", type=Path, default=ROOT / "ThesisDocs/defense/defense_slides.json")
    parser.add_argument("--output", type=Path, default=ROOT / "ThesisDocs/defense/native_content_audit.json")
    parser.add_argument("--numeric-tolerance", type=float, default=1e-14)
    args = parser.parse_args()
    if not math.isfinite(args.numeric_tolerance) or args.numeric_tolerance < 0:
        parser.error("numeric tolerance must be finite and nonnegative")
    report = audit(args.pptx.resolve(), args.source.resolve(), args.numeric_tolerance)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps({key: report[key] for key in ("all_passed", "slide_count", "native_table_count", "native_chart_count", "notes_duration_seconds", "defects")}))
    for chart in report["charts"]:
        print(json.dumps({"chart_slide": chart["slide_number"], "categories": chart["categories"], "series": chart["series"]}))
    print(f"Audit: {args.output}")
    raise SystemExit(0 if report["all_passed"] else 1)


if __name__ == "__main__":
    main()
